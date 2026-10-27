"""Original document fixtures; expected text is authored independently of extraction."""

import io
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import cairosvg
import pymupdf
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def create_samples(directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    samples = []

    def record(name, required, **extra):
        samples.append(
            {
                "file": str(directory / name),
                "required_text": required,
                "outcome": "success",
                **extra,
            }
        )

    simple = pymupdf.open()
    page = simple.new_page()
    page.insert_text((60, 70), "A simple public-style report", fontsize=20)
    page.insert_text((60, 110), "The river level increased steadily during the observation period.")
    simple.save(directory / "simple.pdf")
    simple.close()
    record("simple.pdf", ["river level increased steadily"])

    columns = pymupdf.open()
    for i in range(2):
        page = columns.new_page()
        page.insert_text((50, 20), "REPEATED RUNNING HEADER")
        page.insert_text((50, 80), f"Left column {i + 1}", fontsize=16)
        page.insert_textbox(
            (50, 100, 270, 600),
            "The left column explains the first observation.\n\nIt ends before the right column begins.",
            fontsize=12,
        )
        page.insert_text((320, 80), f"Right column {i + 1}", fontsize=16)
        page.insert_textbox(
            (320, 100, 550, 600),
            "The right column explains the next observation.\n\nIts final sentence completes this page.",
            fontsize=12,
        )
        page.insert_text((50, 820), "REPEATED RUNNING FOOTER")
    columns.save(directory / "columns.pdf")
    columns.close()
    record(
        "columns.pdf",
        ["first observation", "next observation", "final sentence"],
        excluded_text=["REPEATED RUNNING HEADER", "REPEATED RUNNING FOOTER"],
        required_order=["Left column 1", "Right column 1", "Left column 2", "Right column 2"],
    )

    chart = cairosvg.svg2png(url="tests/corpus/assets/chart.svg")
    illustrated = pymupdf.open()
    page = illustrated.new_page()
    page.insert_text((50, 60), "Illustrated measurement report", fontsize=20)
    for x in [50, 220, 390, 550]:
        page.draw_line((x, 95), (x, 175))
    for y in [95, 120, 145, 175]:
        page.draw_line((50, y), (550, y))
    for x, text in [(60, "Location"), (230, "Minimum"), (400, "Maximum")]:
        page.insert_text((x, 112), text)
    for y, values in [(138, ["North", "12", "24"]), (164, ["South", "15", "27"])]:
        for x, text in zip([60, 230, 400], values, strict=True):
            page.insert_text((x, y), text)
    page.insert_image(pymupdf.Rect(50, 220, 550, 470), stream=chart)
    page.insert_text((50, 500), "The chart retains the underlying measured pattern.")
    page = illustrated.new_page()
    page.insert_text((50, 60), "Vector plot", fontsize=18)
    page.draw_polyline(
        [
            pymupdf.Point(50, 300),
            pymupdf.Point(180, 250),
            pymupdf.Point(310, 270),
            pymupdf.Point(500, 100),
        ],
        color=(0, 0, 0),
        width=3,
    )
    page.insert_text((50, 350), "Vector figures must remain visible.")
    illustrated.save(directory / "illustrated.pdf")
    illustrated.close()
    record(
        "illustrated.pdf",
        ["North", "South", "24", "27", "measured pattern"],
        required_tags=["table", "img"],
    )

    docx = Document()
    docx.add_heading("A document with useful structure", 0)
    docx.add_paragraph("Résumé: the original paragraphs remain in order.")
    docx.add_heading("Tasks", 1)
    docx.add_paragraph("Measure the distance.", style="List Number")
    docx.add_paragraph("Record the result.", style="List Number")
    table = docx.add_table(rows=2, cols=2)
    for cell, text in zip(
        [*table.rows[0].cells, *table.rows[1].cells],
        ["Item", "Value", "Distance", "42 m"],
        strict=True,
    ):
        cell.text = text
    docx.add_picture(io.BytesIO(chart))
    docx.save(directory / "structured.docx")
    record(
        "structured.docx",
        ["Résumé", "Measure the distance.", "Record the result.", "42 m"],
        required_tags=["table", "img", "ol"],
    )

    numbered = docx.add_paragraph("A native numbered paragraph.")
    numpr = OxmlElement("w:numPr")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(docx.styles["List Number"].element.pPr.numPr.numId.val))
    numpr.append(numid)
    numbered._p.get_or_add_pPr().append(numpr)
    cell = docx.add_table(rows=2, cols=3)
    cell.cell(0, 0).merge(cell.cell(0, 1)).text = "Shared label"
    cell.cell(0, 2).text = "Third column"
    for index, text in enumerate(["One", "Two", "Three"]):
        cell.cell(1, index).text = text
    paragraph = cell.cell(1, 2).paragraphs[0]
    hyperlink = OxmlElement("w:hyperlink")
    rel = docx.part.relate_to(
        "https://example.com/reference", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = " Reference"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    buffer = io.BytesIO()
    docx.save(buffer)
    with ZipFile(io.BytesIO(buffer.getvalue())) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    word = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    tree = ET.fromstring(files["word/document.xml"])
    paragraph = tree.find(f".//{{{word}}}p")
    run = ET.SubElement(paragraph, f"{{{word}}}r")
    ET.SubElement(run, f"{{{word}}}footnoteReference", {f"{{{word}}}id": "2"})
    files["word/document.xml"] = ET.tostring(tree)
    footnotes = ET.Element(f"{{{word}}}footnotes")
    footnote = ET.SubElement(footnotes, f"{{{word}}}footnote", {f"{{{word}}}id": "2"})
    p = ET.SubElement(footnote, f"{{{word}}}p")
    r = ET.SubElement(p, f"{{{word}}}r")
    ET.SubElement(r, f"{{{word}}}t").text = "A supporting source is retained in this footnote."
    files["word/footnotes.xml"] = ET.tostring(footnotes)
    relpath = "word/_rels/document.xml.rels"
    rels = ET.fromstring(files[relpath])
    ET.SubElement(
        rels,
        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
        {
            "Id": "rIdTestFootnotes",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
            "Target": "footnotes.xml",
        },
    )
    files[relpath] = ET.tostring(rels)
    ct = ET.fromstring(files["[Content_Types].xml"])
    ET.SubElement(
        ct,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
        {
            "PartName": "/word/footnotes.xml",
            "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        },
    )
    files["[Content_Types].xml"] = ET.tostring(ct)
    with ZipFile(directory / "notes-and-tables.docx", "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    record(
        "notes-and-tables.docx",
        ["A supporting source is retained", "Shared label", "A native numbered paragraph."],
        required_tags=["table", "a", "ol"],
        required_selectors=[{"selector": "ol li", "text": "A native numbered paragraph."}],
        required_links=[{"href": "https://example.com/reference", "text": "Reference"}],
    )

    (directory / "unicode.txt").write_text(
        "Reading notes\n\nCafé and 日本語 remain readable.\n\nA second paragraph follows.\n"
    )
    record("unicode.txt", ["Café and 日本語", "second paragraph"])
    (directory / "code.md").write_text(
        "# A small script\n\nRun these steps:\n\n1. Open the file.\n2. Read each line.\n\n```python\nfor line in lines:\n    print(line)\n```\n\n[Reference](https://example.com/reference)\n"
    )
    record("code.md", ["Read each line.", "    print(line)"], required_tags=["pre", "ol", "a"])

    scanned = pymupdf.open()
    scanned.new_page().insert_image(pymupdf.Rect(50, 50, 550, 300), stream=chart)
    scanned.save(directory / "scanned.pdf")
    scanned.close()
    protected = pymupdf.open()
    protected.new_page().insert_text((50, 60), "Protected reading content")
    protected.save(
        directory / "encrypted.pdf",
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        user_pw="fixture-password",
        owner_pw="fixture-owner",
    )
    protected.close()
    (directory / "corrupt.pdf").write_bytes(b"%PDF-1.7\ntruncated-invalid-content")
    (directory / "disguised.pdf").write_text(
        '<html><body><form><input type="password"/>Sign in</form></body></html>'
    )
    (directory / "unsupported.bin").write_bytes(b"\x00\x00\x01\x02unsupported binary")
    for name, reason in [
        ("scanned.pdf", "OCR"),
        ("encrypted.pdf", "encrypted"),
        ("corrupt.pdf", "PDF"),
        ("disguised.pdf", "PDF"),
        ("unsupported.bin", "Unsupported"),
    ]:
        samples.append({"file": str(directory / name), "outcome": "held", "reason": reason})
    return samples


if __name__ == "__main__":
    import json

    samples = create_samples("output/corpus/documents")
    Path("output/corpus/documents/manifest.json").write_text(json.dumps(samples, indent=2) + "\n")
    print(f"Created {len(samples)} document cases")
